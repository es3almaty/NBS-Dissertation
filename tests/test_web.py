from __future__ import annotations

import io
import os
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

os.environ.setdefault("RESULT_TTL_MINUTES", "5")

from fastapi.testclient import TestClient
from docx import Document

from app.main import app, store


def make_minimal_docx(path: Path, *, complete: bool) -> None:
    doc = Document()
    doc.add_paragraph("NARXOZ UNIVERSITY")
    doc.add_paragraph("MASTER'S PROJECT")
    doc.add_paragraph("Student Name")
    if complete:
        doc.add_heading("CONTENTS", 1)
        doc.add_paragraph("1. Main body")
        doc.add_heading("PROJECT SUMMARY", 1)
        doc.add_paragraph("Summary text")
        doc.add_heading("АННОТАЦИЯ", 1)
        doc.add_paragraph("Қазақша аннотация мәтіні")
        doc.add_paragraph("Русская аннотация")
        doc.add_paragraph("English annotation")
    doc.add_heading("1. MAIN BODY", 1)
    doc.add_paragraph("Body text with (Smith, 2024).")
    doc.add_heading("CONCLUSION", 1)
    doc.add_paragraph("Conclusion text")
    if complete:
        doc.add_heading("REFERENCES", 1)
        doc.add_paragraph("Smith, J. (2024). Example source. https://doi.org/10.1000/example")
    doc.save(path)


def test_health_and_home(tmp_path):
    with TestClient(app) as client:
        assert client.get("/health").json()["status"] == "ok"
        r = client.get("/?lang=ru")
        assert r.status_code == 200
        assert "Проверьте магистерскую работу" in r.text


def test_invalid_docx_is_rejected():
    with TestClient(app) as client:
        r = client.post(
            "/process",
            data={"language": "en"},
            files={"dissertation": ("fake.docx", b"not a docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            follow_redirects=False,
        )
        assert r.status_code == 400


def test_resubmit_does_not_offer_formatted(tmp_path):
    source = tmp_path / "incomplete.docx"
    make_minimal_docx(source, complete=False)
    with TestClient(app) as client:
        with source.open("rb") as f:
            r = client.post(
                "/process",
                data={"language": "en"},
                files={"dissertation": (source.name, f.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                follow_redirects=False,
            )
        assert r.status_code == 303
        result = client.get(r.headers["location"])
        assert result.status_code == 200
        assert "RESUBMIT BEFORE FORMATTING" in result.text
        assert "Formatted dissertation" not in result.text
        assert "Non-compliance report" in result.text


def test_ready_or_conditional_route_offers_report_and_formatted(tmp_path):
    source = tmp_path / "complete.docx"
    make_minimal_docx(source, complete=True)
    with TestClient(app) as client:
        with source.open("rb") as f:
            r = client.post(
                "/process",
                data={"language": "en"},
                files={"dissertation": (source.name, f.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                follow_redirects=False,
            )
        assert r.status_code == 303
        location = r.headers["location"]
        page = client.get(location)
        assert page.status_code == 200
        assert "Formatted dissertation" in page.text
        assert "Compliance report" in page.text
        job_id = location.split("/result/")[1].split("?")[0]
        d1 = client.get(f"/download/{job_id}/formatted")
        d2 = client.get(f"/download/{job_id}/report")
        assert d1.status_code == 200 and d1.content[:2] == b"PK"
        assert d2.status_code == 200 and d2.content[:2] == b"PK"
        deleted = client.post(f"/delete/{job_id}", data={"language": "en"}, follow_redirects=False)
        assert deleted.status_code == 303
        assert client.get(f"/result/{job_id}?lang=en").status_code == 410
