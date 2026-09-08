from __future__ import annotations
import zipfile
from pathlib import Path
import pytest
from dissertation_formatter.engine import audit_document
from dissertation_formatter.models import SubmissionStatus, TableRisk, CitationSystem
from dissertation_formatter.ooxml_utils import count_page_fields_by_part
from .fixtures import *


def codes(a):return {f.code for f in a.findings}


def test_01_fully_compliant_ready(tmp_path):
    p=build_base(tmp_path/'full.docx')
    a=audit_document(p,tmp_path/'out','en')
    assert a.status==SubmissionStatus.READY
    assert a.missing_critical_count==0
    assert Path(a.formatted_path).exists() and Path(a.report_path).exists()


def test_02_exactly_two_missing_is_conditional_and_formats(tmp_path):
    p=build_base(tmp_path/'two_missing.docx',contents=False,summary=False)
    a=audit_document(p,tmp_path/'out','en')
    assert a.missing_critical_count==2
    assert a.status==SubmissionStatus.CONDITIONAL
    assert a.formatted_path and Path(a.formatted_path).exists()


def test_03_three_missing_is_resubmit_and_no_format(tmp_path):
    p=build_base(tmp_path/'three_missing.docx',contents=False,summary=False,annotations=False)
    a=audit_document(p,tmp_path/'out','en')
    assert a.missing_critical_count==3
    assert a.status==SubmissionStatus.RESUBMIT
    assert a.formatted_path is None
    assert 'NONCOMPLIANCE_REPORT' in a.report_path


def test_04_bibliography_present_unmatched_citation(tmp_path):
    p=build_base(tmp_path/'unmatched.docx',citation='(Jones, 2025)')
    a=audit_document(p,tmp_path/'out','en')
    assert a.references.references_present
    assert 'UNMATCHED_CITATIONS' in codes(a)


def test_05_numeric_citations_no_bibliography(tmp_path):
    p=build_base(tmp_path/'numeric.docx',references=False,citation='[1] and [4, p. 45] and [25]')
    a=audit_document(p,tmp_path/'out','en')
    assert a.references.system==CitationSystem.NUMERIC
    assert a.references.numeric_citations==[1,4,25]
    assert 'NUMERIC_CITATIONS_NO_BIBLIOGRAPHY' in codes(a)


def test_06_duplicate_table_numbering(tmp_path):
    p=build_base(tmp_path/'duptable.docx');add_duplicate_table_numbers(p)
    a=audit_document(p,tmp_path/'out','en')
    assert 'DUPLICATE_TABLE_NUMBERS' in codes(a)


def test_07_broken_table(tmp_path):
    p=build_base(tmp_path/'broken.docx');add_broken_table(p)
    a=audit_document(p,tmp_path/'out','en')
    assert any(t.risk==TableRisk.BROKEN_REQUIRES_REVIEW for t in a.tables)
    assert 'BROKEN_TABLE' in codes(a)


def test_08_duplicate_page_fields_removed_on_format(tmp_path):
    p=build_base(tmp_path/'pages.docx');add_duplicate_page_fields(p)
    a=audit_document(p,tmp_path/'out','en')
    assert a.preflight.duplicate_page_fields
    assert 'DUPLICATE_PAGE_FIELDS' in codes(a)
    counts=count_page_fields_by_part(a.formatted_path)
    assert all(v<=1 for k,v in counts.items() if 'header' in k or 'footer' in k)


def test_09_automatic_heading_numbering_detected(tmp_path):
    p=build_base(tmp_path/'autonum.docx');add_automatic_heading_numbering(p)
    a=audit_document(p,tmp_path/'out','en')
    assert a.preflight.automatic_numbering_paragraphs
    assert 'AUTOMATIC_WORD_NUMBERING_DETECTED' in codes(a)
    # The generated number is stored in Word numbering XML, not visible text, but
    # the heading must still count as a major numbered section.
    assert a.metadata['major_section_count']==4
    assert any(x.get('number_source')=='word_numbering' for x in a.metadata['major_sections'])


def test_10_contents_mismatch(tmp_path):
    p=build_base(tmp_path/'tocbad.docx',contents_mismatch=True)
    a=audit_document(p,tmp_path/'out','en')
    assert 'CONTENTS_HEADING_MISMATCH' in codes(a)


def test_11_table_prose_numerical_inconsistency(tmp_path):
    p=build_base(tmp_path/'numdiff.docx');add_numeric_inconsistency(p)
    a=audit_document(p,tmp_path/'out','en')
    fs=[f for f in a.findings if f.code=='TABLE_PROSE_NUMERIC_INCONSISTENCY']
    assert fs
    assert any(f.evidence.get('table_value')=='6.43' and f.evidence.get('prose_value')=='6.13' for f in fs)


def test_12_appendix_crossref_mismatch(tmp_path):
    p=build_base(tmp_path/'appbad.docx');add_appendix_reference_without_appendix(p)
    a=audit_document(p,tmp_path/'out','en')
    assert 'APPENDIX_REFERENCE_MISMATCH' in codes(a)
    assert a.components['appendices'].status.value=='MISSING'


def test_13_numbered_lists_do_not_become_major_sections(tmp_path):
    p=build_base(tmp_path/'listnoise.docx');add_numbered_lists_that_are_not_major_sections(p)
    a=audit_document(p,tmp_path/'out','en')
    assert a.metadata['major_section_count']==3
    assert [x['number'] for x in a.metadata['major_sections']]==[1,2,3]
    assert 'SECTION_NUMBERING_GAPS' not in codes(a)


def test_14_caption_association_skips_blanks_and_joins_wrap(tmp_path):
    p=build_base(tmp_path/'captionwrap.docx');add_caption_with_blank_and_wrap(p)
    a=audit_document(p,tmp_path/'out','en')
    assert 'DUPLICATE_TABLE_NUMBERS' in codes(a)
    caps=[t.caption for t in a.tables if t.number==10]
    assert any(c and 'в швейной компании' in c for c in caps)


def test_15_narrow_index_column_is_not_broken(tmp_path):
    p=build_base(tmp_path/'indexcol.docx');add_narrow_index_column_table(p)
    a=audit_document(p,tmp_path/'out','en')
    target=[t for t in a.tables if t.number==20][0]
    assert target.risk!=TableRisk.BROKEN_REQUIRES_REVIEW


def test_16_malformed_figure_caption_still_detects_duplicate(tmp_path):
    p=build_base(tmp_path/'figartefact.docx');add_malformed_duplicate_figure_caption(p)
    a=audit_document(p,tmp_path/'out','en')
    assert 'DUPLICATE_FIGURE_NUMBERS' in codes(a)


def test_17_repeated_metric_value_inconsistency(tmp_path):
    p=build_base(tmp_path/'metricdiff.docx');add_repeated_metric_inconsistency(p)
    a=audit_document(p,tmp_path/'out','en')
    fs=[f for f in a.findings if f.code=='REPEATED_METRIC_VALUE_INCONSISTENCY']
    assert fs
    assert set(fs[0].evidence['values']) >= {1.72,4.4,5.0}


def test_18_unused_footnotes_part_does_not_count_as_footnote_system(tmp_path):
    p=build_base(tmp_path/'unusednotes.docx');add_unused_footnotes_part(p)
    a=audit_document(p,tmp_path/'out','en')
    assert a.preflight.footnotes_present is False
    assert a.references.system!=CitationSystem.FOOTNOTE


def test_19_takim_obrazom_counts_as_section_conclusion(tmp_path):
    p=build_base(tmp_path/'synthesis.docx');replace_section_conclusion_markers_with_takim_obrazom(p)
    a=audit_document(p,tmp_path/'out','en')
    assert 'SECTION_CONCLUSIONS_MISSING' not in codes(a)


def test_20_post_format_content_integrity_is_verified(tmp_path):
    p=build_base(tmp_path/'integrity.docx')
    a=audit_document(p,tmp_path/'out','en')
    assert a.metadata['post_format_content_integrity']['passed'] is True
    assert 'POST_FORMAT_CONTENT_INTEGRITY_VERIFIED' in codes(a)

def test_21_integrity_guard_detects_text_tampering(tmp_path):
    from dissertation_formatter.integrity_guard import verify_content_integrity
    from dissertation_formatter.word_formatter import format_document
    from dissertation_formatter.models import CitationSystem
    from docx import Document
    p=build_base(tmp_path/'source.docx')
    out=tmp_path/'formatted.docx'
    format_document(p,out,CitationSystem.APA_AUTHOR_DATE)
    d=Document(out)
    target=next(x for x in d.paragraphs if x.text.startswith('Проект анализирует'))
    target.text=target.text+' ИЗМЕНЕНО'
    d.save(out)
    r=verify_content_integrity(p,out)
    assert r.passed is False
    assert 'document' in r.failures

def test_22_static_toc_is_flagged_for_manual_page_update(tmp_path):
    p=build_base(tmp_path/'static_toc.docx')
    remove_dynamic_toc_field(p)
    a=audit_document(p,tmp_path/'out','en')
    assert 'STATIC_TOC_PAGE_NUMBERS_REQUIRE_UPDATE' in codes(a)
    assert a.status==SubmissionStatus.CONDITIONAL

def test_23_title_page_table_formatting_is_preserved(tmp_path):
    from dissertation_formatter.word_formatter import format_document
    from dissertation_formatter.models import CitationSystem
    from docx import Document
    p=build_title_table_fixture(tmp_path/'title_table.docx')
    before=Document(p)
    bp=before.tables[0].cell(0,1).paragraphs[0]
    # Fixture inherits normal spacing/font: formatter must not impose body double-spacing
    # or 12-pt direct formatting on the title-page supervisor table.
    out=tmp_path/'title_table_formatted.docx'
    format_document(p,out,CitationSystem.NO_DETECTABLE_SYSTEM)
    after=Document(out)
    ap=after.tables[0].cell(0,1).paragraphs[0]
    assert ap.paragraph_format.line_spacing==bp.paragraph_format.line_spacing
    assert ap.runs[0].font.size==bp.runs[0].font.size

def test_24_engine_never_releases_output_when_content_changes(tmp_path,monkeypatch):
    import dissertation_formatter.engine as eng
    from docx import Document
    from dissertation_formatter.integrity_guard import ContentIntegrityError
    p=build_base(tmp_path/'source.docx')
    real=eng.format_document
    def corrupting_formatter(src,dst,citation_system):
        fixes=real(src,dst,citation_system)
        d=Document(dst)
        target=next(x for x in d.paragraphs if x.text.startswith('Проект анализирует'))
        target.text=target.text+' CORRUPTED'
        d.save(dst)
        return fixes
    monkeypatch.setattr(eng,'format_document',corrupting_formatter)
    with pytest.raises(ContentIntegrityError):
        eng.audit_document(p,tmp_path/'out','en')
    assert not list((tmp_path/'out').glob('*FORMATTED_DISSERTATION.docx'))
