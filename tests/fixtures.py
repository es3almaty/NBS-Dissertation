from __future__ import annotations

import zipfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _page_break(doc):
    p=doc.add_paragraph();p.add_run().add_break(WD_BREAK.PAGE)


def _add_title(doc):
    for line in [
        'НАО «УНИВЕРСИТЕТ НАРХОЗ»',
        'Иванов Иван Иванович',
        'СТРАТЕГИЯ ПОВЫШЕНИЯ ЭФФЕКТИВНОСТИ КОМПАНИИ',
        'Код и наименование ОП: 7M04101 MBA',
        'МАГИСТЕРСКИЙ ПРОЕКТ',
        'на соискание степени Магистра делового администрирования',
        'Научный руководитель: Петров П.П., PhD',
        'Алматы, 2026',
    ]:
        p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _page_break(doc)


def _add_contents(doc, mismatch=False):
    doc.add_paragraph('СОДЕРЖАНИЕ')
    doc.add_paragraph(('1 Неверное название' if mismatch else '1 Исследовательская часть') + ' ........ 5')
    doc.add_paragraph('2 Результаты исследования ........ 12')
    doc.add_paragraph('3 Практические рекомендации ........ 20')
    doc.add_paragraph('ЗАКЛЮЧЕНИЕ ........ 30')
    doc.add_paragraph('СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ ........ 32')
    _page_break(doc)


def _add_project_summary(doc):
    doc.add_paragraph('РЕЗЮМЕ ПРОЕКТА')
    doc.add_paragraph('Актуальность темы определяется практической задачей. Практическая значимость состоит в применимости рекомендаций. Цель проекта состоит в анализе процесса; задачи включают сбор и анализ данных; объект исследования - организация; практическая база - данные организации.')
    _page_break(doc)


def _add_annotations(doc):
    doc.add_paragraph('Қазақша аннотация')
    doc.add_paragraph('Жобаның қысқаша сипаттамасы және негізгі нәтижелері берілген.')
    doc.add_paragraph('Аннотация на русском языке')
    doc.add_paragraph('Краткая характеристика проекта и его основных результатов.')
    doc.add_paragraph('Annotation in English')
    doc.add_paragraph('Brief description of the project and its principal results.')
    _page_break(doc)


def _add_main(doc, citation='(Smith, 2024)'):
    doc.add_paragraph('ВВЕДЕНИЕ')
    doc.add_paragraph('Проект анализирует практическую управленческую проблему. ' + citation)
    for n,title in [(1,'Исследовательская часть'),(2,'Результаты исследования'),(3,'Практические рекомендации')]:
        doc.add_paragraph(f'{n} {title}')
        doc.add_paragraph(f'{n}.1 Подраздел {n}')
        doc.add_paragraph('Основной текст раздела содержит достаточно материала для проверки структуры и сохраняет академическое содержание без автоматического переписывания.')
        doc.add_paragraph(f'Выводы по разделу {n}')
        doc.add_paragraph('Краткие выводы данного раздела сформулированы студентом.')
    doc.add_paragraph('ЗАКЛЮЧЕНИЕ')
    doc.add_paragraph('В заключении кратко представлены результаты и общие рекомендации проекта.')


def _add_references(doc, entries=None):
    doc.add_paragraph('СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ')
    for e in entries or ['Smith, J. (2024). Example research source. Academic Press.']:
        doc.add_paragraph(e)


def build_base(path: Path, *, contents=True, summary=True, annotations=True, references=True, contents_mismatch=False, citation='(Smith, 2024)', ref_entries=None):
    doc=Document();_add_title(doc)
    if contents:_add_contents(doc,contents_mismatch)
    if summary:_add_project_summary(doc)
    if annotations:_add_annotations(doc)
    _add_main(doc,citation)
    if references:_add_references(doc,ref_entries)
    doc.save(path)
    return path


def add_duplicate_table_numbers(path: Path):
    doc=Document(path)
    for x in range(2):
        doc.add_paragraph('Таблица 1 – Тестовая таблица')
        t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='Показатель';t.cell(0,1).text='Значение';t.cell(1,0).text='A';t.cell(1,1).text=str(x+1)
    doc.save(path)


def add_numeric_inconsistency(path: Path):
    doc=Document(path)
    doc.add_paragraph('Таблица 1 – Показатели деловой активности')
    t=doc.add_table(rows=2,cols=4)
    for j,v in enumerate(['Показатель','2021','2022','2023']):t.cell(0,j).text=v
    for j,v in enumerate(['Оборачиваемость запасов','2,77','6,43','3,78']):t.cell(1,j).text=v
    doc.add_paragraph('Оборачиваемость запасов достигла пикового значения в 2022 году и составила 6,13, после чего снизилась.')
    doc.save(path)


def add_appendix_reference_without_appendix(path: Path):
    doc=Document(path);doc.add_paragraph('Дополнительные расчеты приведены в Приложение 2.');doc.save(path)


def add_broken_table(path: Path):
    doc=Document(path);doc.add_paragraph('Таблица 1 – Основные показатели платежеспособности')
    t=doc.add_table(rows=2,cols=5)
    vals=['Показатель','Коэффициент текущей ликвидности','Коэффициент быстрой ликвидности','Коэффициент абсолютной ликвидности','Норматив']
    for i,v in enumerate(vals):t.cell(0,i).text=v
    for i,v in enumerate(['2023','1.2','0.8','0.3','1.0']):t.cell(1,i).text=v
    doc.save(path)
    _patch_table_grid(path,250)


def _patch_table_grid(path: Path,width: int):
    tmp=path.with_suffix('.tmp.docx')
    with zipfile.ZipFile(path,'r') as zin,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename=='word/document.xml':
                from lxml import etree
                root=etree.fromstring(data);ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for gc in root.xpath('.//w:tbl[last()]/w:tblGrid/w:gridCol',namespaces=ns):
                    gc.set('{%s}w'%ns['w'],str(width))
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
            zout.writestr(item,data)
    tmp.replace(path)


def add_duplicate_page_fields(path: Path):
    doc=Document(path);p=doc.sections[0].header.paragraphs[0]
    for _ in range(2):
        fs=OxmlElement('w:fldSimple');fs.set(qn('w:instr'),'PAGE');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);fs.append(r);p._p.append(fs)
    doc.save(path)


def add_automatic_heading_numbering(path: Path):
    doc=Document(path)
    heading_text='Автоматически нумерованный раздел'
    p=doc.add_paragraph(heading_text,style='Heading 1')
    # Place it before the real Conclusion heading so it is a main-body candidate rather
    # than post-reference material. The fixture is Russian, so do not assume an English
    # literal. The visible paragraph text deliberately has no generated number.
    conclusion_labels={'заключение','conclusion','қорытынды'}
    target=next(x for x in doc.paragraphs if x.text.strip().lower() in conclusion_labels)
    target._p.addprevious(p._p)
    doc.save(path)

    # Attach an existing decimal numbering definition to THIS heading.  The previous
    # fixture patched the final body paragraph, which could accidentally number a
    # reference entry rather than the generated Heading 1.
    with zipfile.ZipFile(path) as zf:
        styles=zf.read('word/styles.xml')
    from lxml import etree
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    root=etree.fromstring(styles);numid='5'
    for st in root.xpath('.//w:style[@w:styleId="ListNumber"]',namespaces=ns):
        vals=st.xpath('./w:pPr/w:numPr/w:numId/@w:val',namespaces=ns)
        if vals:
            numid=vals[0]
            break

    tmp=path.with_suffix('.tmp.docx')
    with zipfile.ZipFile(path,'r') as zin,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename=='word/document.xml':
                root=etree.fromstring(data)
                matches=[]
                for para in root.xpath('.//w:body/w:p',namespaces=ns):
                    text=''.join(para.xpath('.//w:t/text()',namespaces=ns)).strip()
                    if text==heading_text:
                        matches.append(para)
                if len(matches)!=1:
                    raise AssertionError(f'Automatic-numbering fixture expected exactly one target heading, found {len(matches)}')
                target_el=matches[0]
                pPr=target_el.find('{%s}pPr'%ns['w'])
                if pPr is None:
                    pPr=etree.Element('{%s}pPr'%ns['w']);target_el.insert(0,pPr)
                # Replace any inherited/direct numPr so the test is deterministic.
                existing=pPr.find('{%s}numPr'%ns['w'])
                if existing is not None:
                    pPr.remove(existing)
                numPr=etree.Element('{%s}numPr'%ns['w'])
                ilvl=etree.SubElement(numPr,'{%s}ilvl'%ns['w']);ilvl.set('{%s}val'%ns['w'],'0')
                nid=etree.SubElement(numPr,'{%s}numId'%ns['w']);nid.set('{%s}val'%ns['w'],numid)
                pPr.insert(0,numPr)
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
            zout.writestr(item,data)
    tmp.replace(path)


def add_numbered_lists_that_are_not_major_sections(path: Path):
    """Add ordinary top-level-looking numbered list items; section count must stay 3."""
    doc=Document(path)
    conclusion=next(p for p in doc.paragraphs if p.text.strip().lower()=='заключение')
    for text in [
        '8. Влияние изменения себестоимости продукции:',
        '9. Влияние изменения количества сотрудников на итоговую прибыль.',
        '10. Влияние изменения заработной платы на итоговую прибыль:',
        '12. Для оценки изменения при сохранении других факторов используется дополнительный расчет.',
    ]:
        p=doc.add_paragraph(text)
        conclusion._p.addprevious(p._p)
    doc.save(path)


def add_caption_with_blank_and_wrap(path: Path):
    doc=Document(path)
    doc.add_paragraph('Таблица 10 – График предоставления бюджетных заявок ЦФО')
    doc.add_paragraph('в швейной компании')
    doc.add_paragraph('')
    t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='Действие';t.cell(0,1).text='Срок';t.cell(1,0).text='Подготовка';t.cell(1,1).text='Октябрь'
    doc.add_paragraph('Таблица 10 – Другая таблица')
    doc.add_paragraph('')
    t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='A';t.cell(0,1).text='B';t.cell(1,0).text='1';t.cell(1,1).text='2'
    doc.save(path)


def add_narrow_index_column_table(path: Path):
    doc=Document(path);doc.add_paragraph('Таблица 20 – Таблица с узкой колонкой номера')
    t=doc.add_table(rows=3,cols=3)
    vals=[['№','Категория','Описание'],['1','Планирование','Очень длинное описание процесса, которое находится в широкой колонке и не означает, что узкая колонка номера повреждена.'],['2','Контроль','Еще одно достаточно длинное описание для проверки эвристики.']]
    for r,row in enumerate(vals):
        for c,v in enumerate(row): t.cell(r,c).text=v
    doc.save(path)
    # Make only the first grid column narrow; this is a legitimate index column.
    tmp=path.with_suffix('.tmp.docx')
    with zipfile.ZipFile(path,'r') as zin,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename=='word/document.xml':
                from lxml import etree
                ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                root=etree.fromstring(data); tbl=root.xpath('.//w:tbl[last()]',namespaces=ns)[0]
                cols=tbl.xpath('./w:tblGrid/w:gridCol',namespaces=ns)
                cols[0].set('{%s}w'%ns['w'],'420');cols[1].set('{%s}w'%ns['w'],'2500');cols[2].set('{%s}w'%ns['w'],'6500')
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
            zout.writestr(item,data)
    tmp.replace(path)


def add_malformed_duplicate_figure_caption(path: Path):
    doc=Document(path)
    doc.add_paragraph('Рисунок 15 - Влияние результатов анализа')
    doc.add_paragraph('Рисуноко 15 -о Общаяо схемао управленияо бизнес-процессом')
    doc.save(path)


def add_repeated_metric_inconsistency(path: Path):
    doc=Document(path)
    doc.add_paragraph('Срок окупаемости составляет 1,72 года – это расчетный срок окупаемости проекта.')
    doc.add_paragraph('Окупаемость составляет 5 лет или 4,4 года при точном расчете.')
    doc.add_paragraph('Итоговый срок окупаемости проекта составляет около 1,72 года.')
    doc.save(path)


def add_unused_footnotes_part(path: Path):
    """Add a footnotes.xml part but no w:footnoteReference in document.xml."""
    tmp=path.with_suffix('.tmp.docx')
    xml=b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotes>'''
    with zipfile.ZipFile(path,'r') as zin,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        names={x.filename for x in zin.infolist()}
        for item in zin.infolist(): zout.writestr(item,zin.read(item.filename))
        if 'word/footnotes.xml' not in names: zout.writestr('word/footnotes.xml',xml)
    tmp.replace(path)


def replace_section_conclusion_markers_with_takim_obrazom(path: Path):
    doc=Document(path)
    for p in doc.paragraphs:
        if p.text.startswith('Выводы по разделу'):
            p.text='Таким образом, результаты раздела обобщены и сформулированы.'
    doc.save(path)
