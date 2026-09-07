from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from .models import ComponentResult, ComponentStatus, Finding, Severity, Action
from .ooxml_utils import paragraph_numbering_info


def norm(s: str) -> str:
    s = s.replace("ё", "е").lower().strip()
    s = re.sub(r"[\s\u00a0]+", " ", s)
    s = re.sub(r"[.:;\-–—]+$", "", s).strip()
    return s

HEADING_PATTERNS = {
    "contents": [r"^содержание$", r"^оглавление$", r"^contents$", r"^мазмұны$"],
    "project_summary": [r"^резюме проекта$", r"^project summary$", r"^жоба резюмесі$", r"^жобаның резюмесі$", r"^жоба түйіндемесі$", r"^жобаның түйіндемесі$"],
    "conclusion": [r"^заключение$", r"^conclusion$", r"^қорытынды$"],
    "references": [r"^список использованной литературы$", r"^список литературы$", r"^библиография$", r"^references$", r"^reference list$", r"^пайдаланылған әдебиеттер$", r"^әдебиеттер тізімі$"],
    "appendix": [r"^приложение\s+\d+[а-яa-z-]*$", r"^appendix\s+\d+[a-z-]*$", r"^қосымша\s+\d+[а-яa-z-]*$"],
}

ANNOTATION_PATTERNS = {
    "ru": [r"^аннотация(?:\s+на\s+русском(?:\s+языке)?)?$", r"^русская аннотация$"],
    "en": [r"^annotation(?:\s+in\s+english)?$", r"^abstract(?:\s+in\s+english)?$", r"^english annotation$"],
    "kk": [r"^аннотация(?:\s+на\s+казахском(?:\s+языке)?)?$", r"^қазақша аннотация$", r"^аннотация\s*\(қазақша\)$", r"^аңдатпа$"],
}

TITLE_REQUIRED_SIGNAL_GROUPS = [
    ("narxoz", ["университет нархоз", "narxoz university", "нархоз университеті", "narxoz университеті"]),
    ("masters_project", ["магистерский проект", "master's project", "masters project", "магистрлік жоба"]),
    ("supervisor", ["научный руководитель", "supervisor", "research supervisor", "ғылыми жетекші"]),
    ("almaty", ["алматы", "almaty"]),
]


def _find_heading(paragraphs: list[str], patterns: list[str]) -> list[int]:
    idxs = []
    for i, text in enumerate(paragraphs):
        n = norm(text)
        if any(re.match(p, n, re.I) for p in patterns):
            idxs.append(i)
    return idxs


def detect_components(path: str | Path) -> tuple[dict[str, ComponentResult], list[Finding], dict]:
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs]
    text = "\n".join(paras)
    low = norm(text)
    findings: list[Finding] = []
    metadata: dict = {}

    # Title page: first 25 paragraphs, separate presence from template compliance.
    first = norm("\n".join(paras[:25]))
    title_presence = bool(first.strip()) and ("магист" in first or "master" in first or "нархоз" in first or "narxoz" in first)
    title_hits = [label for label,variants in TITLE_REQUIRED_SIGNAL_GROUPS if any(v in first for v in variants)]
    first_raw="\n".join(paras[:15])
    has_name_like = bool(
        re.search(r"[А-ЯӘҒҚҢӨҰҮҺЁ][а-яәғқңөұүһё]+\s+[А-ЯӘҒҚҢӨҰҮҺЁ][а-яәғқңөұүһё]+", first_raw) or
        re.search(r"\b[A-Z][A-Za-z'’-]{2,}\s+[A-Z][A-Za-z'’-]{2,}\b", first_raw)
    )
    has_year = bool(re.search(r"\b20\d{2}\b", "\n".join(paras[:25])))
    has_program = bool(re.search(r"\b(?:mba|emba|op|оп|бб|образовательн\w* программ\w*|educational program|degree program|program)\b", first, re.I))
    title_compliant = title_presence and len(title_hits) >= 3 and has_name_like and has_year and has_program
    title_status = ComponentStatus.PRESENT_COMPLIANT if title_compliant else ComponentStatus.PRESENT_NONCOMPLIANT if title_presence else ComponentStatus.MISSING

    heads = {k: _find_heading(paras, pats) for k, pats in HEADING_PATTERNS.items()}

    def cr(name, status, evidence=None, notes=None, applicable=True):
        return ComponentResult(name=name, status=status, critical=True, applicable=applicable, evidence=evidence or [], notes=notes or [])

    comps: dict[str, ComponentResult] = {
        "title_page": cr("Title page", title_status, [f"Template signals found: {', '.join(title_hits) or 'none'}"], []),
        "contents": cr("Contents", ComponentStatus.PRESENT_COMPLIANT if heads["contents"] else ComponentStatus.MISSING, [f"Heading at paragraph {heads['contents'][0]+1}"] if heads["contents"] else []),
        "project_summary": cr("Project Summary", ComponentStatus.PRESENT_COMPLIANT if heads["project_summary"] else ComponentStatus.MISSING, [f"Heading at paragraph {heads['project_summary'][0]+1}"] if heads["project_summary"] else []),
        "main_body": cr("Main body", ComponentStatus.PRESENT_COMPLIANT if _has_main_body(paras) else ComponentStatus.MISSING),
        "conclusion": cr("Conclusion", ComponentStatus.PRESENT_COMPLIANT if heads["conclusion"] else ComponentStatus.MISSING, [f"Heading at paragraph {heads['conclusion'][0]+1}"] if heads["conclusion"] else []),
        "references": cr("References", ComponentStatus.PRESENT_COMPLIANT if heads["references"] else ComponentStatus.MISSING, [f"Heading at paragraph {heads['references'][0]+1}"] if heads["references"] else []),
    }

    anno = {}
    for lang, pats in ANNOTATION_PATTERNS.items():
        anno[lang] = _find_heading(paras, pats)
    anno_count = sum(bool(v) for v in anno.values())
    if anno_count == 3:
        astat = ComponentStatus.PRESENT_COMPLIANT
    elif anno_count == 0:
        astat = ComponentStatus.MISSING
    else:
        astat = ComponentStatus.PRESENT_NONCOMPLIANT
    comps["annotations_three_languages"] = cr(
        "Three-language annotation set", astat,
        [f"Detected annotation languages: {', '.join(k for k,v in anno.items() if v) or 'none'}"],
        ["Required set: Kazakh, Russian, English"]
    )

    appendix_refs = _appendix_references(text)
    appendix_heads = heads["appendix"]
    appendix_required = bool(appendix_refs or appendix_heads)
    if appendix_required and appendix_heads:
        app_status = ComponentStatus.PRESENT_COMPLIANT
    elif appendix_required:
        app_status = ComponentStatus.MISSING
    else:
        app_status = ComponentStatus.PRESENT_COMPLIANT
    comps["appendices"] = cr(
        "Appendices", app_status,
        [f"Appendix references detected: {sorted(appendix_refs)}", f"Appendix headings detected: {len(appendix_heads)}"],
        ["Critical only when appendices are cited or clearly required."],
        applicable=appendix_required,
    )

    if title_status == ComponentStatus.PRESENT_NONCOMPLIANT:
        findings.append(Finding("TITLE_TEMPLATE_NONCOMPLIANT", "Title page present but non-compliant", "A title page is present, but it does not contain enough signals from the Regulation Appendix 2 template to be treated as compliant.", Severity.MAJOR, Action.FLAG))

    # Three sections is advisory, not blocking because Regulation says 'as a rule'.
    majors = detect_major_sections(paras, path=path, doc=doc)
    metadata["major_section_count"] = len(majors)
    metadata["major_sections"] = majors
    if majors and len(majors) != 3:
        findings.append(Finding(
            "NORMAL_THREE_SECTION_STRUCTURE", "Structure differs from normal three-section form",
            f"Detected {len(majors)} major numbered sections. The Regulation says the main part should, as a rule, consist of three sections. Confirm this structure with the supervisor/program.",
            Severity.ADVISORY, Action.MANUAL_REVIEW,
        ))
    # Section conclusions.
    if majors:
        missing_conclusions = _major_sections_without_conclusions(paras, majors)
        if missing_conclusions:
            findings.append(Finding(
                "SECTION_CONCLUSIONS_MISSING", "Section conclusions may be missing",
                "The Regulation states that each main section should end with conclusions. No conclusion marker was detected near the end of: " + ", ".join(str(x) for x in missing_conclusions),
                Severity.MAJOR, Action.FLAG,
            ))
    return comps, findings, metadata


def _has_main_body(paras: list[str]) -> bool:
    meaningful = [p for p in paras if len(p.split()) >= 5]
    if len(meaningful) < 8:
        return False
    return bool(detect_major_sections(paras) or any(norm(p) in {"введение", "introduction", "кіріспе"} for p in paras))


def _upper_ratio(text: str) -> float:
    letters=[c for c in text if c.isalpha()]
    return (sum(c.isupper() for c in letters)/len(letters)) if letters else 0.0


def _followed_by_first_subsection(paras: list[str], index: int, number: int, window: int=12) -> bool:
    pat=re.compile(rf"^{number}\.1\.?\s+\S")
    for p in paras[index+1:min(len(paras),index+1+window)]:
        t=p.strip()
        if not t:
            continue
        if pat.match(t):
            return True
        # Do not scan through another apparent major-number paragraph.
        if re.match(r"^\d+\.?\s+\S",t) and not re.match(r"^\d+\.\d+",t):
            break
    return False


def detect_major_sections(paras: list[str], path: str|Path|None=None, doc=None) -> list[dict]:
    """Detect level-1 dissertation sections conservatively.

    Student documents contain many ordinary numbered lists. A visible paragraph such as
    "8. Влияние изменения ..." is therefore *not* enough to establish a chapter. For a
    Normal-style visible number we require a structural corroborator: the expected N.1
    subsection nearby, heading/outline styling, or strongly heading-like capitalization.
    Word-generated Heading 1 numbering remains a first-class structural signal.
    """
    if doc is None and path is not None:
        try:
            doc=Document(path)
        except Exception:
            doc=None

    # Main-body candidates stop at Conclusion/References.
    stop=len(paras)
    for i,p in enumerate(paras):
        if norm(p) in {"заключение","conclusion","қорытынды","список использованной литературы","список литературы","references","пайдаланылған әдебиеттер","әдебиеттер тізімі"}:
            stop=min(stop,i)

    explicit=[]
    for i,p in enumerate(paras[:stop]):
        t=p.strip()
        if re.search(r"\.{2,}\s*\d+\s*$",t):
            continue
        m=re.match(r"^(\d+)\.?\s+(.{3,})$",t)
        if not m or re.match(r"^\d+\.\d+",t):
            continue
        n=int(m.group(1)); title=m.group(2).strip()
        if len(title)>180:
            continue
        style_name=""
        outline_level=None
        if doc is not None and i < len(doc.paragraphs):
            para=doc.paragraphs[i]
            style_name=((para.style.name if para.style else "") or "").lower()
            try:
                ol=para._p.pPr.outlineLvl if para._p.pPr is not None else None
                outline_level=int(ol.val) if ol is not None else None
            except Exception:
                outline_level=None
        style_heading=style_name.startswith("heading 1") or style_name.startswith("заголовок 1") or outline_level==0
        subsection_signal=_followed_by_first_subsection(paras,i,n)
        capitalization_signal=len(title)>=10 and _upper_ratio(title)>=0.72
        if style_heading or subsection_signal or capitalization_signal:
            explicit.append({
                "number":n,"index":i,"title":title,"number_source":"visible_text",
                "signals":{"heading_style":style_heading,"subsection":subsection_signal,"capitalization":capitalization_signal},
            })

    auto=[]
    if path is not None:
        try:
            if doc is None:
                doc=Document(path)
            for d in paragraph_numbering_info(path):
                i=d.get("paragraph_index")
                if i is None or i>=stop or i>=len(paras) or str(d.get("ilvl","0"))!="0":
                    continue
                text=paras[i].strip()
                if not text or len(text)>180 or re.match(r"^\d+(?:\.\d+)*\s+",text):
                    continue
                style_name=((doc.paragraphs[i].style.name if doc.paragraphs[i].style else "") or "").lower()
                style_id=(d.get("style") or "").lower()
                is_heading1=style_name in {"heading 1","заголовок 1"} or style_name.startswith("heading 1") or style_name.startswith("заголовок 1") or style_id in {"heading1","1"}
                if is_heading1:
                    auto.append({"number":None,"index":i,"title":text,"number_source":"word_numbering","numId":d.get("numId"),"ilvl":d.get("ilvl")})
        except Exception:
            auto=[]

    by_index={d["index"]:d for d in explicit}
    for d in auto:
        by_index.setdefault(d["index"],d)
    merged=[by_index[i] for i in sorted(by_index)]

    # Assign a conservative ordinal to Word-generated numbers whose display text is not
    # materialized. This is only for structural count/sequence checking.
    last=0; used={d["number"] for d in merged if d.get("number") is not None}
    for d in merged:
        if d.get("number") is not None:
            last=d["number"]
        else:
            cand=last+1
            while cand in used:
                cand+=1
            d["number"]=cand;used.add(cand);last=cand

    # Keep the earliest credible heading for an exact number collision. Do not infer
    # chapters merely because numbered list items later use larger integers.
    best={}
    for d in merged:
        best.setdefault(d["number"],d)
    return [best[k] for k in sorted(best)]

def _major_sections_without_conclusions(paras: list[str], majors: list[dict]) -> list[int]:
    missing=[]
    # The general dissertation Conclusion is not a substitute for section conclusions.
    # For the final main section, stop before the global Conclusion/References heading.
    global_stop=len(paras)
    for i,p in enumerate(paras):
        if norm(p) in {"заключение","conclusion","қорытынды","список использованной литературы","список литературы","references","пайдаланылған әдебиеттер","әдебиеттер тізімі"}:
            global_stop=min(global_stop,i)
    for j, d in enumerate(majors):
        start=d["index"]
        end=majors[j+1]["index"] if j+1 < len(majors) else global_stop
        nonempty=[p for p in paras[start:end] if p.strip()]
        tail="\n".join(nonempty[-20:]).lower()
        # Students often signal a genuine section synthesis with "Таким образом" rather
        # than a literal heading "Выводы". Treat that as a conclusion marker.
        if not re.search(r"(?:вывод|итог|таким\s+образом|conclusion|қорытынды)", tail, re.I):
            missing.append(d["number"])
    return missing


def _appendix_references(text: str) -> set[int]:
    refs=set()
    for m in re.finditer(r"(?:приложени[еяю]|appendix|қосымша)\s*(?:№\s*)?(\d+)", text, re.I):
        refs.add(int(m.group(1)))
    return refs
