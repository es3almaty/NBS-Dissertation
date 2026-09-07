from __future__ import annotations

import re
from pathlib import Path
from collections import Counter
from docx import Document
from .models import Finding, Severity, Action
from .structure_detector import detect_major_sections, norm
from .table_engine import CAPTION_RE, normalize_caption_text

# Tolerate the specific class of stray leading-caption characters seen in real Word
# documents without normalizing arbitrary prose.
FIG_RE = re.compile(r"^(?:рисунок(?:о)?|рис\.?|figure|сурет)\s*(\d+)\b\s*[-–—:.]?\s*(?:о\s+)?(.*)$", re.I)
SECTION_REF_RE = re.compile(r"(?:раздел(?:е|а|у)?|section)\s+(\d+(?:\.\d+)*)", re.I)
TABLE_REF_RE = re.compile(r"(?:таблиц(?:а|е|ы|у)|table)\s+(\d+)", re.I)
FIG_REF_RE = re.compile(r"(?:рисун(?:ок|ке|ка)|рис\.?|figure)\s+(\d+)", re.I)
APP_REF_RE = re.compile(r"(?:приложени[еяю]|appendix|қосымша)\s*(?:№\s*)?(\d+)", re.I)
SOURCE_PLACEHOLDER_RE = re.compile(r"(?:источник|source|дереккөз)\s*[:.-]?\s*\[\s*\]", re.I)
PAYBACK_TERM_RE = re.compile(r"(?:срок\s+окупаемости|окупаемость(?:\s+проекта)?|payback\s+period)", re.I)
YEAR_VALUE_RE = re.compile(r"(-?\d+(?:[.,]\d+)?)\s*(?:года?|лет|years?)\b", re.I)


def _normalize_figure_caption(text: str) -> str:
    s=(text or "").strip()
    s=re.sub(r"^(рисунок)о\s+",r"\1 ",s,flags=re.I)
    s=re.sub(r"^(рисунок|рис\.?|figure|сурет)\s*(\d+)\s*([-–—:.])\s*о\s+",r"\1 \2 \3 ",s,flags=re.I)
    return s


def _iter_body(doc):
    body = doc.element.body
    pidx = -1
    tidx = -1
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            pidx += 1
            texts = child.xpath(".//w:t/text()")
            yield "p", pidx, "".join(texts).strip(), None
        elif child.tag.endswith("}tbl"):
            tidx += 1
            yield "t", tidx, None, doc.tables[tidx]


def _actual_section_numbers(paras):
    nums = set()
    for p in paras:
        m = re.match(r"^(\d+(?:\.\d+)*)\.?\s+", p.strip())
        if m:
            nums.add(m.group(1))
    return nums


def _contents_mismatches(paras):
    start = None
    for i, p in enumerate(paras):
        if norm(p) in {"содержание", "оглавление", "contents", "мазмұны"}:
            start = i
            break
    if start is None:
        return []
    entries = []
    stop_heads = {"резюме проекта", "project summary", "жоба резюмесі", "қазақша аннотация", "аннотация на русском языке", "annotation in english", "введение", "introduction", "кіріспе"}
    for p in paras[start + 1:start + 80]:
        if norm(p) in stop_heads:
            break
        if re.match(r"^\d+\.?\s+\S", p) and not re.search(r"\.{2,}|\s\d{1,3}$", p):
            break
        cleaned = re.sub(r"\.{2,}\s*\d+\s*$", "", p).strip()
        cleaned = re.sub(r"\s+\d+\s*$", "", cleaned).strip()
        if cleaned and len(cleaned) > 3:
            entries.append(cleaned)
    actual = []
    for p in paras:
        if re.match(r"^\d+(?:\.\d+)*\.?\s+\S", p.strip()) or norm(p) in {"введение", "заключение", "список использованной литературы", "references", "conclusion", "introduction"}:
            actual.append(p.strip())
    actual_norm = {norm(re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", a)): a for a in actual}
    mism = []
    for e in entries:
        e2 = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", e)
        if norm(e2) not in actual_norm:
            mism.append(e)
    return mism[:20]


def _table_prose_numeric_mismatches(doc) -> list[dict]:
    """Conservative detector for a table metric/year/value contradicted in nearby prose."""
    body = list(_iter_body(doc))
    out = []
    for pos, (typ, idx, txt, tbl) in enumerate(body):
        if typ != "t" or not tbl.rows:
            continue
        rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
        if not rows or len(rows[0]) < 3:
            continue
        header = rows[0]
        years = {j: cell for j, cell in enumerate(header) if re.fullmatch(r"20\d{2}", cell.strip())}
        if not years:
            continue
        nearby = []
        for typ2, idx2, txt2, t2 in body[pos + 1:pos + 9]:
            if typ2 == "t":
                break
            if txt2:
                nearby.append(txt2)
        prose = " ".join(nearby)
        for r in rows[1:]:
            if len(r) < max(years) + 1:
                continue
            label = norm(r[0])
            words = [w for w in re.findall(r"[a-zа-яәғқңөұүһё]{5,}", label) if w not in {"коэффициент", "показатель", "период", "значение"}]
            if len(label) < 8 or not words:
                continue
            distinctive = words[:2]
            for col, year in years.items():
                cell = r[col].replace(",", ".").strip()
                if not re.fullmatch(r"-?\d+(?:\.\d+)?", cell):
                    continue
                pattern = r"(?i)" + ".*?".join(map(re.escape, distinctive[:1])) + r"[^.]{0,240}?" + re.escape(year) + r"[^.]{0,100}?(-?\d+[.,]\d+)"
                m = re.search(pattern, prose, re.I)
                if not m:
                    continue
                pval = m.group(1).replace(",", ".")
                try:
                    if abs(float(pval) - float(cell)) > 1e-9:
                        out.append({"metric": r[0], "year": year, "table_value": cell, "prose_value": pval})
                except ValueError:
                    pass
    seen = set()
    final = []
    for d in out:
        key = tuple(d.values())
        if key not in seen:
            final.append(d)
            seen.add(key)
    return final


def _cluster_close(values: list[float]) -> list[list[float]]:
    """Group rounded representations such as 1.704 and 1.72 together."""
    clusters: list[list[float]]=[]
    for v in sorted(values):
        placed=False
        for c in clusters:
            ref=sum(c)/len(c)
            tol=max(0.05, abs(ref)*0.02)
            if abs(v-ref)<=tol:
                c.append(v);placed=True;break
        if not placed:
            clusters.append([v])
    return clusters


def _repeated_metric_value_inconsistencies(paras: list[str]) -> list[dict]:
    """Detect incompatible values for a repeated, explicitly named metric.

    v0.2 starts with payback period because it is a mechanically identifiable quantity
    with an explicit unit. The logic is generic across Russian/English phrasing and does
    not decide which value is correct.
    """
    observations=[]
    for i,p in enumerate(paras):
        if not PAYBACK_TERM_RE.search(p):
            continue
        vals=[]
        for m in YEAR_VALUE_RE.finditer(p):
            try: vals.append(float(m.group(1).replace(",",".")))
            except ValueError: pass
        if vals:
            observations.append({"paragraph_index":i,"values":vals,"text":p[:500]})
    all_values=[v for o in observations for v in o["values"] if 0 <= v <= 100]
    clusters=_cluster_close(all_values)
    if len(clusters)<=1:
        return []
    # Only flag when the clusters are materially different; ordinary rounding stays in
    # one cluster. Retain the raw evidence and do not select a preferred value.
    reps=[round(sum(c)/len(c),4) for c in clusters]
    return [{"metric":"payback period","values":sorted(set(round(v,4) for v in all_values)),"clusters":reps,"observations":observations}]


def check_internal_consistency(path: str | Path) -> list[Finding]:
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs]
    text = "\n".join(paras)
    findings = []

    hits = SOURCE_PLACEHOLDER_RE.findall(text)
    if hits:
        findings.append(Finding("EMPTY_SOURCE_PLACEHOLDER", "Empty source placeholder detected", f"Detected {len(hits)} empty source placeholder(s), such as 'источник []'. Student review is required.", Severity.MAJOR, Action.FLAG))

    # Figure caption numbers, with conservative leading-artifact normalization.
    fnums = []
    for p in paras:
        m = FIG_RE.match(_normalize_figure_caption(p.strip()))
        if m:
            fnums.append(int(m.group(1)))
    dup = sorted(n for n, c in Counter(fnums).items() if c > 1)
    if dup:
        findings.append(Finding("DUPLICATE_FIGURE_NUMBERS", "Duplicate figure numbers", "Duplicate figure numbers detected: " + ", ".join(map(str, dup)), Severity.MAJOR, Action.FLAG, evidence={"numbers":dup}))
    if fnums:
        u = sorted(set(fnums))
        g = [n for n in range(min(u), max(u) + 1) if n not in u]
        if g:
            findings.append(Finding("FIGURE_NUMBERING_GAPS", "Figure numbering gaps", "Missing figure numbers: " + ", ".join(map(str, g)), Severity.MINOR, Action.FLAG, evidence={"numbers":g}))

    tcaption = set()
    fcaption = set()
    appheads = set()
    for p in paras:
        m = CAPTION_RE.match(normalize_caption_text(p.strip()))
        if m:
            tcaption.add(int(m.group(1)))
        m = FIG_RE.match(_normalize_figure_caption(p.strip()))
        if m:
            fcaption.add(int(m.group(1)))
        m = re.match(r"^(?:приложение|appendix|қосымша)\s+(\d+)", p.strip(), re.I)
        if m:
            appheads.add(int(m.group(1)))
    trefs = {int(x) for x in TABLE_REF_RE.findall(text)}
    frefs = {int(x) for x in FIG_REF_RE.findall(text)}
    arefs = {int(x) for x in APP_REF_RE.findall(text)}
    missing_t = sorted(trefs - tcaption) if tcaption else sorted(trefs) if trefs else []
    missing_f = sorted(frefs - fcaption) if fcaption else sorted(frefs) if frefs else []
    missing_a = sorted(arefs - appheads)
    if missing_t:
        findings.append(Finding("NONEXISTENT_TABLE_REFERENCE", "References to tables without matching captions", "Table references with no matching detected caption: " + ", ".join(map(str, missing_t)), Severity.MAJOR, Action.FLAG))
    if missing_f:
        findings.append(Finding("NONEXISTENT_FIGURE_REFERENCE", "References to figures without matching captions", "Figure references with no matching detected caption: " + ", ".join(map(str, missing_f)), Severity.MAJOR, Action.FLAG))
    if missing_a:
        findings.append(Finding("APPENDIX_REFERENCE_MISMATCH", "Appendix cross-reference mismatch", "Appendix references with no matching appendix heading: " + ", ".join(map(str, missing_a)), Severity.CRITICAL, Action.BLOCK_FORMATTING))

    actual = _actual_section_numbers(paras)
    srefs = set(SECTION_REF_RE.findall(text))
    missing_s = sorted(srefs - actual)
    if missing_s:
        findings.append(Finding("NONEXISTENT_SECTION_REFERENCE", "Reference to nonexistent section", "Section references with no matching detected section heading: " + ", ".join(missing_s), Severity.MAJOR, Action.FLAG))

    majors = detect_major_sections(paras, path=path, doc=doc)
    nums = [d["number"] for d in majors]
    if nums:
        gaps = [n for n in range(min(nums), max(nums) + 1) if n not in nums]
        if gaps:
            findings.append(Finding("SECTION_NUMBERING_GAPS", "Section numbering gaps", "Missing major section numbers: " + ", ".join(map(str, gaps)), Severity.MAJOR, Action.FLAG))

    cm = _contents_mismatches(paras)
    if cm:
        findings.append(Finding("CONTENTS_HEADING_MISMATCH", "Contents entries differ from detected headings", "Contents entries could not be matched exactly to actual headings: " + "; ".join(cm), Severity.MAJOR, Action.FLAG, evidence={"entries": cm}))

    for d in _table_prose_numeric_mismatches(doc):
        findings.append(Finding(
            "TABLE_PROSE_NUMERIC_INCONSISTENCY", "Table/prose numerical inconsistency",
            f"Internal inconsistency detected: the table reports {d['table_value']} for '{d['metric']}' in {d['year']}, while nearby prose reports {d['prose_value']}. The system cannot determine which value is correct. Please verify.",
            Severity.MAJOR, Action.MANUAL_REVIEW, evidence=d,
        ))

    for d in _repeated_metric_value_inconsistencies(paras):
        vals=", ".join(str(v) for v in d["values"])
        findings.append(Finding(
            "REPEATED_METRIC_VALUE_INCONSISTENCY", "Repeated metric has conflicting values",
            f"Internal inconsistency detected: the same metric ('{d['metric']}') is reported with materially different values in the document ({vals} years). The system cannot determine which value is correct. Please verify the cited calculations and narrative.",
            Severity.MAJOR, Action.MANUAL_REVIEW, evidence=d,
        ))
    return findings
