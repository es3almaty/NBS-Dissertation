from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from .models import CitationSystem
from .ooxml_utils import count_page_fields_by_part, remove_extra_page_fields_from_xml
from .integrity_guard import has_dynamic_toc



_FRONT_MATTER_RX=re.compile(r"^(?:аннотация|аңдатпа|abstract|содержание|contents|мазмұны|резюме проекта|project summary|жоба түйіндемесі)\b",re.I)

def _title_page_end_index(doc: Document) -> int:
    """Return the first paragraph index after the title page."""
    for i,p in enumerate(doc.paragraphs):
        if _FRONT_MATTER_RX.match(p.text.strip()):
            return i
    return 0

def _title_page_tables(doc: Document) -> set[object]:
    """Return XML table elements appearing before the first front-matter heading.

    Supervisor details are often placed in a table on NBS title pages. Those tables must
    be preserved with the rest of the title-page layout rather than globally double-spaced.
    """
    preserved=set()
    for child in doc.element.body.iterchildren():
        if child.tag==qn("w:p"):
            p=Paragraph(child,doc)
            if _FRONT_MATTER_RX.match(p.text.strip()):
                break
        elif child.tag==qn("w:tbl"):
            preserved.add(child)
    return preserved


def _set_update_fields_on_open(path: Path) -> bool:
    """Ask Word to refresh dynamic fields (especially a real TOC) when the file opens."""
    if not has_dynamic_toc(path):
        return False
    tmp=path.with_suffix(".fields.tmp.docx")
    settings_name="word/settings.xml"
    with zipfile.ZipFile(path,"r") as zin, zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename==settings_name:
                from lxml import etree
                root=etree.fromstring(data)
                nodes=root.xpath('./w:updateFields',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
                if nodes:
                    node=nodes[0]
                else:
                    node=etree.SubElement(root,qn('w:updateFields'))
                node.set(qn('w:val'),'true')
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
            zout.writestr(item,data)
    tmp.replace(path)
    return True

def _set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar");tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{m}"))
        if node is None:
            node=OxmlElement(f"w:{m}");tcMar.append(node)
        node.set(qn("w:w"),str(v));node.set(qn("w:type"),"dxa")


def _add_page_field(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run()
    begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "
    separate=OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"),"separate")
    text=OxmlElement("w:t"); text.text="1"
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    run._r.extend([begin,instr,separate,text,end])


def _dedupe_page_fields_in_package(path: Path) -> int:
    counts=count_page_fields_by_part(path)
    targets=[p for p,c in counts.items() if ("header" in p or "footer" in p) and c>1]
    if not targets:return 0
    tmp=path.with_suffix(".tmp.docx");removed=0
    with zipfile.ZipFile(path,"r") as zin, zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename in targets:
                data,r=remove_extra_page_fields_from_xml(data,keep=1);removed+=r
            zout.writestr(item,data)
    tmp.replace(path)
    return removed


def _apply_reference_format(doc: Document, citation_system: CitationSystem):
    start=None
    for i,p in enumerate(doc.paragraphs):
        if p.text.strip().lower() in {"references","список литературы","список использованной литературы","библиография","пайдаланылған әдебиеттер","әдебиеттер тізімі"}:
            start=i; p.paragraph_format.page_break_before=True; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.bold=True
            break
    if start is None:return
    for p in doc.paragraphs[start+1:]:
        if not p.text.strip():continue
        p.paragraph_format.left_indent=Cm(1.27)
        p.paragraph_format.first_line_indent=Cm(-1.27)
        p.paragraph_format.line_spacing=2


def format_document(input_path: str|Path, output_path: str|Path, citation_system: CitationSystem) -> list[str]:
    input_path=Path(input_path);output_path=Path(output_path)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(input_path,output_path)
    doc=Document(output_path);fixes=[]
    # A4, portrait, 2.54cm margins from Narxoz APA guidance.
    for sec in doc.sections:
        sec.orientation=WD_ORIENT.PORTRAIT
        sec.page_width=Cm(21.0);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54);sec.left_margin=Cm(2.54);sec.right_margin=Cm(2.54)
    fixes.append("Set A4 portrait with 2.54 cm margins")

    # Normalize body run typography without touching data/text. Preserve the existing
    # title-page typography and spacing as a single layout-sensitive block.
    title_end=_title_page_end_index(doc)
    title_tables=_title_page_tables(doc) if title_end else set()
    if title_end:
        # Make the title/front-matter boundary explicit. Many student files simulate a
        # title-page break with blank paragraphs, which becomes unstable when later text
        # is reformatted. A page-break-before changes layout only and prevents spill.
        doc.paragraphs[title_end].paragraph_format.page_break_before=True
    for i,p in enumerate(doc.paragraphs):
        if i < title_end:
            continue
        p.paragraph_format.line_spacing=2
        style=(p.style.name if p.style else "").lower()
        # Apply first-line indentation only to prose-like paragraphs. Front matter, headings,
        # captions, TOC entries and list/numbered items are left structurally untouched.
        txt=p.text.strip()
        is_special=(
            style.startswith("heading") or style.startswith("toc") or
            re.match(r"^(таблица|table|кесте|рисунок|figure|сурет|содержание|contents|мазмұны|резюме проекта|project summary|жоба түйіндемесі|заключение|conclusion|қорытынды|список|references|әдебиеттер|приложение|appendix|қосымша)\b",txt,re.I) or
            re.match(r"^(?:\d+(?:\.\d+)*[.)]?|[-•–—])\s+",txt) or
            ("......." in txt)
        )
        if txt and len(txt) >= 40 and not is_special:
            p.paragraph_format.first_line_indent=Cm(1.27)
        for r in p.runs:
            r.font.name="Times New Roman";r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
            r._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
    for table in doc.tables:
        if table._tbl in title_tables:
            continue
        for row in table.rows:
            row._tr.get_or_add_trPr()
            for cell in row.cells:
                _set_cell_margins(cell)
                for p in cell.paragraphs:
                    # Do not create an unsupported table-only typography exception. The Narxoz
                    # guide specifies Times New Roman 12 pt and double spacing for the work.
                    p.paragraph_format.line_spacing=2
                    p.paragraph_format.first_line_indent=None
                    for r in p.runs:
                        r.font.name="Times New Roman";r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
                        r._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
    fixes.append("Normalized body text to Times New Roman 12 pt, black, double-spaced")
    if title_end:
        fixes.append("Preserved existing title-page layout to prevent pagination spill")

    _apply_reference_format(doc,citation_system)
    if citation_system in {CitationSystem.APA_AUTHOR_DATE,CitationSystem.MIXED}:
        fixes.append("Applied 1.27 cm hanging indent to supplied reference entries")

    # Ensure a PAGE field exists in the first section header if no PAGE field is already in package.
    doc.save(output_path)
    removed=_dedupe_page_fields_in_package(output_path)
    if removed:
        fixes.append(f"Removed {removed} duplicate PAGE field(s) from header/footer parts")
    counts=count_page_fields_by_part(output_path)
    if sum(counts.values())==0:
        doc=Document(output_path)
        header=doc.sections[0].header
        p=header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        _add_page_field(p)
        doc.save(output_path)
        fixes.append("Inserted a single right-aligned PAGE field in the header")
    if _set_update_fields_on_open(output_path):
        fixes.append("Configured Word to update the dynamic table of contents/fields on open")
    return fixes
