from __future__ import annotations

import re
import zipfile
from pathlib import Path
from docx import Document
from .models import PreflightResult
from .ooxml_utils import NS, count_page_fields_by_part, paragraph_numbering_info, read_zip_xml, zip_parts, field_instructions


def run_preflight(path: str | Path) -> PreflightResult:
    path = Path(path)
    if path.suffix.lower() != ".docx":
        raise ValueError("Input must be a .docx file")
    if not zipfile.is_zipfile(path):
        raise ValueError("Input is not a valid DOCX/ZIP package")
    doc = Document(path)
    parts = zip_parts(path)
    page_counts = count_page_fields_by_part(path)
    numbered = paragraph_numbering_info(path)
    heading_paras = []
    for i, p in enumerate(doc.paragraphs):
        style = (p.style.name if p.style else "") or ""
        if style.lower().startswith("heading") or re.match(r"^(\d+(?:\.\d+)*)\s+\S", p.text.strip()):
            heading_paras.append({"index": i, "text": p.text.strip(), "style": style})
    docxml = read_zip_xml(path, "word/document.xml")
    img_count = len([p for p in parts if p.startswith("word/media/")])
    hyperlink_count = len(docxml.xpath(".//w:hyperlink", namespaces=NS)) if docxml is not None else 0
    eq = bool(docxml.xpath(".//m:oMath|.//m:oMathPara", namespaces=NS)) if docxml is not None else False
    fields = field_instructions(docxml)
    header_parts = [p for p in parts if re.match(r"word/header\d+\.xml$", p)]
    footer_parts = [p for p in parts if re.match(r"word/footer\d+\.xml$", p)]
    return PreflightResult(
        paragraphs=len(doc.paragraphs),
        tables=len(doc.tables),
        images=img_count,
        hyperlinks=hyperlink_count,
        sections=len(doc.sections),
        headers=len(header_parts),
        footers=len(footer_parts),
        page_fields=sum(page_counts.values()),
        duplicate_page_fields=any(v > 1 for k, v in page_counts.items() if "header" in k or "footer" in k),
        # Word packages may contain footnotes.xml/endnotes.xml even when the main
        # document has no note references. Count actual references, not package parts.
        footnotes_present=bool(docxml.xpath(".//w:footnoteReference", namespaces=NS)) if docxml is not None else False,
        endnotes_present=bool(docxml.xpath(".//w:endnoteReference", namespaces=NS)) if docxml is not None else False,
        equations_present=eq,
        fields_present=bool(fields) or any(v for v in page_counts.values()),
        automatic_numbering_paragraphs=[d["paragraph_index"] for d in numbered],
        heading_paragraphs=heading_paras,
        section_breaks=max(len(doc.sections)-1, 0),
        package_parts=parts,
    )
