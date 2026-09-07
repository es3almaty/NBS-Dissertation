from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from .models import AuditResult, SubmissionStatus, ComponentStatus, Severity
from .translations import (
    tr, rt, display_component, display_component_status,
    display_submission_status, display_table_risk, display_finding_title,
)


def _set_font(run,size=10,bold=False):
    run.font.name="Segoe UI";run.font.size=Pt(size);run.bold=bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"),"Segoe UI")


def _p(doc,text="",size=10,bold=False,space_after=4):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(space_after);p.paragraph_format.line_spacing=1.05
    r=p.add_run(text);_set_font(r,size,bold);return p


def _heading(doc,text,level=1):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(10 if level==1 else 6);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
    r=p.add_run(text);_set_font(r,14 if level==1 else 11,True);return p


def _shade(cell,fill="D9EAF7"):
    tcPr=cell._tc.get_or_add_tcPr();shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd");tcPr.append(shd)
    shd.set(qn("w:fill"),fill)


def _has_cyrillic(s:str)->bool:
    return bool(re.search(r"[А-Яа-яЁёӘәҒғҚқҢңӨөҰұҮүҺһІі]", s or ""))


def _tail_after_colon(s:str)->str:
    return s.split(":",1)[1].strip() if ":" in s else ""


def _localized_message(lang, finding):
    """Localize deterministic engine messages without changing their substance.

    Evidence-based/manual findings that are already written in the selected language are
    preserved verbatim. Technical rule codes remain visible for auditability.
    """
    msg=finding.message or ""
    if lang=='en':
        return msg
    if lang=='ru' and _has_cyrillic(msg):
        return msg
    code=finding.code; ev=finding.evidence or {}
    if lang=='ru':
        if code=='TITLE_TEMPLATE_NONCOMPLIANT': return 'Титульный лист обнаружен, но в нем недостаточно признаков шаблона Приложения 2 Положения, чтобы считать его соответствующим требованиям.'
        if code=='NORMAL_THREE_SECTION_STRUCTURE':
            m=re.search(r'Detected\s+(\d+)',msg); n=m.group(1) if m else '?'
            return f'Обнаружено основных нумерованных разделов: {n}. Положение указывает, что основная часть, как правило, состоит из трех разделов. Структуру следует подтвердить у руководителя/программы; автоматически перестраивать работу нельзя.'
        if code=='SECTION_CONCLUSIONS_MISSING': return 'Положение требует завершать каждый основной раздел выводами. Для одного или нескольких разделов маркер выводов возле конца раздела не обнаружен: '+_tail_after_colon(msg)
        if code=='NUMERIC_CITATIONS_NO_BIBLIOGRAPHY':
            nums=ev.get('numeric_citations') or re.findall(r'\d+',_tail_after_colon(msg))
            return 'Обнаружены числовые внутритекстовые ссылки, но раздел списка литературы отсутствует. Восстанавливать или угадывать библиографию нельзя. Обнаруженные номера: '+', '.join(map(str,nums))
        if code=='UNMATCHED_CITATIONS': return 'Следующие внутритекстовые ссылки не удалось надежно сопоставить с предоставленными записями списка литературы: '+_tail_after_colon(msg)
        if code=='DUPLICATE_REFERENCE_ENTRIES': return f"Обнаружены возможные дубли библиографических записей: {len(ev.get('entries',[])) or '?'} шт."
        if code=='BROKEN_TABLE': return 'Структура/читабельность таблицы выглядит поврежденной. Система сохраняет таблицу без агрессивного переформатирования и требует ручной проверки. '+_tail_after_colon(msg)
        if code=='DUPLICATE_TABLE_NUMBERS': return 'Обнаружены повторяющиеся номера таблиц: '+_tail_after_colon(msg)
        if code=='TABLE_NUMBERING_GAPS': return 'В последовательности подписей таблиц отсутствуют номера: '+_tail_after_colon(msg)
        if code=='DUPLICATE_FIGURE_NUMBERS': return 'Обнаружены повторяющиеся номера рисунков: '+_tail_after_colon(msg)
        if code=='FIGURE_NUMBERING_GAPS': return 'В последовательности рисунков отсутствуют номера: '+_tail_after_colon(msg)
        if code=='NONEXISTENT_TABLE_REFERENCE': return 'В тексте есть ссылки на таблицы, для которых не обнаружена соответствующая подпись: '+_tail_after_colon(msg)
        if code=='NONEXISTENT_FIGURE_REFERENCE': return 'В тексте есть ссылки на рисунки, для которых не обнаружена соответствующая подпись: '+_tail_after_colon(msg)
        if code=='APPENDIX_REFERENCE_MISMATCH': return 'В тексте есть ссылки на приложения, для которых не обнаружен соответствующий заголовок приложения: '+_tail_after_colon(msg)
        if code=='NONEXISTENT_SECTION_REFERENCE': return 'В тексте есть ссылки на разделы, для которых не обнаружен соответствующий заголовок: '+_tail_after_colon(msg)
        if code=='SECTION_NUMBERING_GAPS': return 'В нумерации основных разделов обнаружены пропуски: '+_tail_after_colon(msg)
        if code=='CONTENTS_HEADING_MISMATCH':
            vals=ev.get('entries') or []
            return 'Некоторые записи в Содержании не удалось точно сопоставить с фактическими заголовками: '+('; '.join(vals) if vals else _tail_after_colon(msg))
        if code=='EMPTY_SOURCE_PLACEHOLDER':
            m=re.search(r'Detected\s+(\d+)',msg); n=m.group(1) if m else '?'
            return f'Обнаружено пустых заполнителей источника (например, «источник []»): {n}. Требуется проверка студентом.'
        if code=='TABLE_PROSE_NUMERIC_INCONSISTENCY':
            if ev:
                return f"Внутреннее несоответствие: таблица сообщает {ev.get('table_value')} для показателя «{ev.get('metric')}» за {ev.get('year')} год, а соседний текст сообщает {ev.get('prose_value')}. Система не может определить правильное значение. Требуется проверка студентом."
            return 'Обнаружено числовое несоответствие между таблицей и сопровождающим текстом. Система не выбирает правильное значение; требуется проверка студентом.'
        if code=='REPEATED_METRIC_VALUE_INCONSISTENCY':
            vals=ev.get('values') or []
            shown=', '.join(str(x).replace('.',',') for x in vals) if vals else '?'
            return f'Один и тот же показатель (срок окупаемости) указан в документе с существенно различающимися значениями: {shown} года/лет. Система не определяет правильное значение. Необходимо проверить расчеты и сопровождающий текст.'
        if code=='DUPLICATE_PAGE_FIELDS': return 'В одном или нескольких колонтитулах обнаружено более одного поля PAGE. При разрешенном форматировании система может безопасно оставить одно поле и удалить дубли.'
        if code=='AUTOMATIC_WORD_NUMBERING_DETECTED':
            m=re.search(r'on\s+(\d+)\s+paragraph',msg); n=m.group(1) if m else '?'
            return f'Обнаружена метаинформация автоматической нумерации Word (numId/стилевая нумерация) в абзацах: {n}. Такие заголовки не считаются ненумерованными только потому, что номер может генерироваться Word.'
        if code=='PAGE_FIELD_LOCATION_NONCOMPLIANT': return 'Поля PAGE обнаружены только в нижнем колонтитуле. Руководство Нархоз для студенческих работ предусматривает номер страницы справа в верхнем колонтитуле. v0.2 сохраняет такой случай и требует ручной проверки, если связи секций неоднозначны.'
        return msg
    # Kazakh output is an engineering translation for beta; preserve evidence if not safely mapped.
    if lang=='kk':
        simple={
            'TITLE_TEMPLATE_NONCOMPLIANT':'Титулдық бет анықталды, бірақ Ереженің 2-қосымшасындағы үлгі белгілері толық емес.',
            'NUMERIC_CITATIONS_NO_BIBLIOGRAPHY':'Мәтінде сандық сілтемелер бар, бірақ әдебиеттер тізімі анықталмады. Библиографияны қайта құрастыруға немесе болжауға болмайды.',
            'TABLE_PROSE_NUMERIC_INCONSISTENCY':'Кесте мен оған жақын мәтінде бір көрсеткіш үшін әртүрлі сандық мәндер анықталды. Жүйе дұрыс мәнді таңдамайды; студент тексеруі қажет.',
            'BROKEN_TABLE':'Кестенің құрылымы/оқылымдылығы бұзылған сияқты. Жүйе оны агрессивті қайта форматтамай сақтайды және қолмен тексеруді талап етеді.',
            'DUPLICATE_PAGE_FIELDS':'Колонтитулда қайталанатын PAGE өрістері анықталды.',
            'AUTOMATIC_WORD_NUMBERING_DETECTED':'Word автоматты нөмірлеу метадеректері (numId/стильдік нөмірлеу) анықталды.'
        }
        return simple.get(code,msg)
    return msg


def _localized_evidence(lang, evidence:list[str], notes:list[str])->str:
    # Evidence may include literal matched headings/indices. Do not rewrite document evidence.
    raw='; '.join(evidence+notes)[:800]
    if not raw: return ''
    if lang=='en': return raw
    # Translate only stable detector scaffolding while preserving literal payloads.
    repl={
        'Template signals found:':'Обнаруженные признаки шаблона:' if lang=='ru' else 'Үлгі белгілері:',
        'Heading at paragraph':'Заголовок в абзаце' if lang=='ru' else 'Тақырып абзацта',
        'Detected annotation languages:':'Обнаруженные языки аннотаций:' if lang=='ru' else 'Анықталған аннотация тілдері:',
        'Required set: Kazakh, Russian, English':'Обязательный набор: казахский, русский, английский' if lang=='ru' else 'Міндетті жиынтық: қазақ, орыс, ағылшын',
        'Appendix references detected:':'Обнаруженные ссылки на приложения:' if lang=='ru' else 'Қосымша сілтемелері:',
        'Appendix headings detected:':'Обнаружено заголовков приложений:' if lang=='ru' else 'Қосымша тақырыптары:',
        'Critical only when appendices are cited or clearly required.':'Критический компонент только когда приложения цитируются или явно требуются.' if lang=='ru' else 'Қосымшаларға сілтеме жасалғанда немесе олар анық қажет болғанда ғана сыни компонент.'
    }
    for a,b in repl.items(): raw=raw.replace(a,b)
    return raw



def _localized_table_reasons(lang, reasons:list[str])->str:
    if not reasons:
        return rt(lang,'no_risk')
    out=[]
    for reason in reasons:
        r=reason
        if lang=='ru':
            if r=='merged cells': r='объединенные ячейки'
            else:
                m=re.fullmatch(r'(\d+) columns',r)
                if m: r=f"{m.group(1)} столбцов"
                m=re.fullmatch(r'(\d+) rows',r)
                if m: r=f"{m.group(1)} строк"
        elif lang=='kk':
            if r=='merged cells': r='біріктірілген ұяшықтар'
            else:
                m=re.fullmatch(r'(\d+) columns',r)
                if m: r=f"{m.group(1)} баған"
                m=re.fullmatch(r'(\d+) rows',r)
                if m: r=f"{m.group(1)} жол"
        out.append(r)
    return '; '.join(out)

def generate_report(audit: AuditResult, output_path: str|Path, source_availability_note: str|None=None) -> Path:
    lang=audit.language if audit.language in {"en","ru","kk"} else "en"
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)

    title={SubmissionStatus.RESUBMIT:tr(lang,"resubmit_title"),SubmissionStatus.READY:tr(lang,"ready_title"),SubmissionStatus.CONDITIONAL:tr(lang,"conditional_title")}[audit.status]
    p=_p(doc,title,20,True,8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=_p(doc,rt(lang,'subtitle'),12,True,12);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _p(doc,f"{rt(lang,'file')}: {Path(audit.input_path).name}",10)
    _p(doc,f"{rt(lang,'status_label')}: {display_submission_status(lang,audit.status.value)}",11,True)
    _p(doc,f"{rt(lang,'missing')}: {audit.missing_critical_count}",10,True)
    if audit.status==SubmissionStatus.RESUBMIT:_p(doc,tr(lang,"resubmit_notice"),11,True,10)
    _p(doc,tr(lang,"source_note"),9,False,8)
    if source_availability_note:_p(doc,source_availability_note,9,True,8)
    doc.add_page_break()

    _heading(doc,tr(lang,"summary"))
    summary_key='provisional_summary' if audit.metadata.get('binary_ooxml_executed') is False else 'engine_summary'
    _p(doc,rt(lang,summary_key,missing=audit.missing_critical_count,major=sum(1 for f in audit.findings if f.severity in {Severity.CRITICAL,Severity.MAJOR}),tables=len(audit.tables),citation=audit.references.system.value))

    _heading(doc,tr(lang,"critical"))
    table=doc.add_table(rows=1,cols=4);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.style="Table Grid"
    hdr=table.rows[0].cells
    for c,t in zip(hdr,[rt(lang,'component'),rt(lang,'component_status'),rt(lang,'applicable'),rt(lang,'evidence')]):
        _shade(c);r=c.paragraphs[0].add_run(t);_set_font(r,9,True)
    for c in audit.components.values():
        row=table.add_row().cells
        vals=[display_component(lang,c.name),display_component_status(lang,c.status.value),rt(lang,'yes') if c.applicable else rt(lang,'no'),_localized_evidence(lang,c.evidence,c.notes)]
        for cell,val in zip(row,vals):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;r=cell.paragraphs[0].add_run(str(val));_set_font(r,8.5)
    _p(doc,"")

    _heading(doc,tr(lang,"corrected"))
    if audit.auto_fixes:
        for x in audit.auto_fixes:_p(doc,"• "+x)
    else:_p(doc,rt(lang,'no_auto'))

    _heading(doc,tr(lang,"actions"))
    action_findings=[f for f in audit.findings if f.severity in {Severity.CRITICAL,Severity.MAJOR}]
    if not action_findings:_p(doc,rt(lang,'no_major'))
    for f in action_findings:
        lt=display_finding_title(lang,f.code,f.title);_p(doc,f"{f.code}: {lt}",10,True,2);_p(doc,_localized_message(lang,f),9.5)

    _heading(doc,tr(lang,"unalterable"))
    manual=[f for f in audit.findings if f.action.value in {"MANUAL_REVIEW","BLOCK_FORMATTING"}]
    if not manual:_p(doc,rt(lang,'none'))
    for f in manual:_p(doc,f"• {display_finding_title(lang,f.code,f.title)}: {_localized_message(lang,f)}",9.5)

    _heading(doc,tr(lang,"internal"))
    internal_codes={"TABLE_PROSE_NUMERIC_INCONSISTENCY","CONTENTS_HEADING_MISMATCH","NONEXISTENT_SECTION_REFERENCE","NONEXISTENT_TABLE_REFERENCE","NONEXISTENT_FIGURE_REFERENCE","APPENDIX_REFERENCE_MISMATCH","EMPTY_SOURCE_PLACEHOLDER","SECTION_NUMBERING_GAPS","REPEATED_METRIC_VALUE_INCONSISTENCY"}
    internal=[f for f in audit.findings if f.code in internal_codes]
    if not internal:_p(doc,rt(lang,'no_internal'))
    for f in internal:_p(doc,"• "+_localized_message(lang,f),9.5)

    _heading(doc,tr(lang,"tables"))
    if not audit.tables:
        _p(doc,rt(lang,'table_check_not_executed') if audit.metadata.get('binary_ooxml_executed') is False else rt(lang,'no_tables'))
    else:
        for t in audit.tables:
            cap=t.caption or (f"Word table index {t.index}" if lang=='en' else f"Таблица Word № {t.index}" if lang=='ru' else f"Word кестесі № {t.index}")
            reasons=_localized_table_reasons(lang,t.reasons)
            _p(doc,f"• {cap}: {display_table_risk(lang,t.risk.value)}. {reasons}",9.2)
    figfind=[f for f in audit.findings if "FIGURE" in f.code or "TABLE" in f.code]
    for f in figfind:_p(doc,f"• {display_finding_title(lang,f.code,f.title)}: {_localized_message(lang,f)}",9.2)

    _heading(doc,tr(lang,"refs"))
    _p(doc,rt(lang,'citation_line',citation=audit.references.system.value,present=rt(lang,'yes') if audit.references.references_present else rt(lang,'no')))
    if audit.references.numeric_citations:_p(doc,rt(lang,'numeric_line',numbers=', '.join(map(str,audit.references.numeric_citations))))
    for f in audit.findings:
        if f.code in {"NUMERIC_CITATIONS_NO_BIBLIOGRAPHY","UNMATCHED_CITATIONS","DUPLICATE_REFERENCE_ENTRIES"}:_p(doc,"• "+_localized_message(lang,f),9.5)
    _p(doc,tr(lang,"reference_disclaimer"),9.5,False,8)

    _heading(doc,tr(lang,"length"))
    _p(doc,rt(lang,'length_note'))
    if audit.metadata.get("major_section_count") is not None:_p(doc,rt(lang,'sections_note',count=audit.metadata.get('major_section_count')))

    _heading(doc,tr(lang,"not_checked"))
    _p(doc,tr(lang,"integrity"));_p(doc,rt(lang,'not_checked_more'))

    _heading(doc,tr(lang,"responsibility"))
    _p(doc,rt(lang,'student_resp'))

    if audit.status==SubmissionStatus.RESUBMIT:
        _heading(doc,tr(lang,"checklist"))
        missing=[display_component(lang,c.name) for c in audit.components.values() if c.applicable and c.critical and c.status==ComponentStatus.MISSING]
        for x in missing:_p(doc,"☐ "+rt(lang,'add_component',component=x))
        for f in action_findings:_p(doc,"☐ "+rt(lang,'review',title=display_finding_title(lang,f.code,f.title)))
        _p(doc,"☐ "+rt(lang,'reupload'))

    _p(doc,rt(lang,'technical_note'),8.5,False,8)
    for s in doc.sections:
        fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer='Narxoz Dissertation Formatter v0.2 | generated compliance report' if lang=='en' else 'Narxoz Dissertation Formatter v0.2 | автоматический отчет о соответствии' if lang=='ru' else 'Narxoz Dissertation Formatter v0.2 | сәйкестік есебі'
        r=fp.add_run(footer);_set_font(r,8)
    output_path=Path(output_path);output_path.parent.mkdir(parents=True,exist_ok=True);doc.save(output_path);return output_path
